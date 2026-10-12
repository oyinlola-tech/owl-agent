from __future__ import annotations

import io

from ..documentResult_document import library_unavailable_result


def extract_word_document(data: bytes, max_chars: int) -> dict:
    try:
        import docx
    except ImportError:
        return library_unavailable_result("docx", "python-docx")
    document = docx.Document(io.BytesIO(data))
    lines = [paragraph.text for paragraph in document.paragraphs if paragraph.text]
    tables = []
    for table in document.tables:
        rows = 0
        cols = 0
        for row in table.rows:
            values = [cell.text for cell in row.cells]
            cols = max(cols, len(values))
            lines.append("\t".join(values))
            rows += 1
        tables.append({"rows": rows, "cols": cols})
    body = "\n".join(lines)
    truncated = False
    if len(body) > max_chars:
        body = body[:max_chars]
        truncated = True
    return {
        "type": "docx",
        "text": body,
        "structure": {"kind": "docx", "paragraphs": len(document.paragraphs), "tables": tables},
        "truncated": truncated,
    }
