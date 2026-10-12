from __future__ import annotations

import os

import docx
import openpyxl


def build_minimal_pdf(text: str) -> bytes:
    content = f"BT /F1 24 Tf 72 720 Td ({text}) Tj ET".encode("latin-1")
    objects = [
        b"<< /Type /Catalog /Pages 2 0 R >>",
        b"<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
        b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] /Contents 4 0 R "
        b"/Resources << /Font << /F1 5 0 R >> >> >>",
        b"<< /Length %d >>\nstream\n%s\nendstream" % (len(content), content),
        b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>",
    ]
    pdf = bytearray(b"%PDF-1.4\n")
    offsets = []
    for index, obj in enumerate(objects, start=1):
        offsets.append(len(pdf))
        pdf += b"%d 0 obj\n%s\nendobj\n" % (index, obj)
    xref_pos = len(pdf)
    pdf += b"xref\n0 %d\n" % (len(objects) + 1)
    pdf += b"0000000000 65535 f \n"
    for off in offsets:
        pdf += b"%010d 00000 n \n" % off
    pdf += b"trailer\n<< /Size %d /Root 1 0 R >>\n" % (len(objects) + 1)
    pdf += b"startxref\n%d\n%%%%EOF" % xref_pos
    return bytes(pdf)


def write_pdf(path: str, text: str) -> None:
    with open(path, "wb") as handle:
        handle.write(build_minimal_pdf(text))


def write_docx(path: str, paragraph: str, table_rows: list[list[str]]) -> None:
    document = docx.Document()
    document.add_paragraph(paragraph)
    table = document.add_table(rows=len(table_rows), cols=len(table_rows[0]))
    for row_index, values in enumerate(table_rows):
        for col_index, value in enumerate(values):
            table.rows[row_index].cells[col_index].text = value
    document.save(path)


def write_xlsx(path: str, sheets: dict[str, list[list]]) -> None:
    workbook = openpyxl.Workbook()
    workbook.remove(workbook.active)
    for name, rows in sheets.items():
        worksheet = workbook.create_sheet(name)
        for row in rows:
            worksheet.append(row)
    workbook.save(path)


def write_text_fixture(path: str, text: str) -> None:
    os.makedirs(os.path.dirname(path), exist_ok=True)
    with open(path, "w", encoding="utf-8") as handle:
        handle.write(text)


def build_documents_project(folder: str) -> str:
    os.makedirs(folder, exist_ok=True)
    write_pdf(os.path.join(folder, "report.pdf"), "Hello OWL document")
    write_docx(
        os.path.join(folder, "memo.docx"),
        "A plain memo paragraph.",
        [["Name", "Role"], ["Ada", "Engineer"]],
    )
    write_xlsx(
        os.path.join(folder, "book.xlsx"),
        {"Alpha": [["h1", "h2"], [1, 2]], "Beta": [["only"], ["beta-cell"]]},
    )
    write_text_fixture(os.path.join(folder, "data.csv"), "name,role\nAda,Engineer\nGrace,Admiral\n")
    write_text_fixture(os.path.join(folder, "notes.txt"), "first line\nsecond line\n")
    return os.path.realpath(folder)
